"""Re-extract 12L12D measurements and verify or rebuild the frozen paper targets."""
import argparse
import json
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "hypocotyl/data"
GENOTYPES = ["Col-0", "hy5", "MLB", "EMS57", "phyAB"]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, help="Rebuild CSVs in a new directory; default verifies bundled targets")
    args = parser.parse_args()
    workbook = load_workbook(DATA / "raw/hypocotyl_growth_20230403.xlsx", data_only=False)
    sheet = workbook["12L12D"]
    rows = []
    for block, hour in enumerate(range(0, 73, 12)):
        first = 2 + 5 * block
        assert sheet.cell(1, first).value == f"{hour}hr"
        for offset, genotype in enumerate(GENOTYPES):
            column = first + offset
            assert sheet.cell(2, column).value == genotype
            for row in range(3, 45):
                cell = sheet.cell(row, column)
                if cell.value is None:
                    continue
                if cell.data_type == "f" or not isinstance(cell.value, (int, float)):
                    raise ValueError(f"Unexpected measurement: {cell.coordinate}")
                rows.append(dict(observation_id=f"12L12D:{cell.coordinate}", sheet="12L12D", genotype=genotype,
                                 elapsed_hours=hour, length=float(cell.value), source_cell=cell.coordinate,
                                 source_row=row, source_column=column))
    observations = pd.DataFrame(rows)
    assignments = pd.read_csv(DATA / "split_assignments.csv")
    assert observations.observation_id.is_unique and assignments.observation_id.is_unique
    assert len(observations) == 943 and set(observations.observation_id) == set(assignments.observation_id)
    observations = observations.merge(assignments, on="observation_id", validate="one_to_one")
    assert observations.groupby(["genotype", "source_row"]).split.nunique().max() == 1
    assert np.isfinite(observations.length).all() and (observations.length > 0).all()
    frozen = DATA / "processed/single_condition_20260909/replicate"
    config = json.loads((frozen / "config.json").read_text())
    if args.output:
        args.output.mkdir(parents=True, exist_ok=False)
    for split in ("train", "val", "test"):
        part = observations[observations.split.eq(split)].copy()
        expected = pd.read_csv(frozen / f"{split}.csv")
        assert set(part.observation_id) == set(expected.observation_id)
        # Restore the saved row order to preserve model target ordering exactly.
        part = part.set_index("observation_id").loc[expected.observation_id].reset_index()
        for column in ("genotype", "elapsed_hours", "length", "source_cell", "source_row", "source_column"):
            np.testing.assert_array_equal(part[column].to_numpy(), expected[column].to_numpy())
        means = part.groupby(["genotype", "elapsed_hours"], sort=True).length.agg(
            mean_length_mm="mean", n_observations="size").reset_index()
        expected_means = pd.read_csv(frozen / f"{split}_means.csv")
        pd.testing.assert_frame_equal(means, expected_means, check_exact=False, rtol=1e-12)
        assert len(part) == config["counts"][split] and len(means) == 35
        if split == "train":
            assert float(part.length.max()) == config["height_scale_mm"]
        if args.output:
            part[expected.columns].to_csv(args.output / f"{split}.csv", index=False)
            means.to_csv(args.output / f"{split}_means.csv", index=False)
    if args.output:
        (args.output / "config.json").write_text(json.dumps(config, indent=2) + "\n")
    print("Verified 943 original measurements, disjoint row groups, 547/171/225 splits, and 35 means per split.")


if __name__ == "__main__":
    main()
