"""Longitudinal hypocotyl forecasting from a masked 0--36 h phenotype prefix."""
from pathlib import Path
import hashlib
import json
import numpy as np
import pandas as pd
import torch
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
DATA=ROOT/'hypocotyl/data'
PREPARED=DATA/'processed/four_genotypes_20260915'
GENOTYPES=['Col-0','hy5','MLB','phyAB']
HOURS=np.arange(0.,73.,3.)
PREFIX_HOURS=np.array([0.,12.,24.,36.])
PROTOCOL='four_genotypes_20260915'


def sha(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def prepare(workbook=None,output=PREPARED):
    workbook=Path(workbook or DATA/'raw/hypocotyl_growth_20230403.xlsx')
    sheet=load_workbook(workbook,data_only=False)['12L12D']
    rows=[]
    for block,hour in enumerate(range(0,73,12)):
        assert sheet.cell(1,2+5*block).value==f'{hour}hr'
        columns={sheet.cell(2,c).value:c for c in range(2+5*block,7+5*block) if sheet.cell(2,c).value is not None}
        if set(columns)!=set(GENOTYPES):
            raise ValueError('Workbook contains a genotype outside the approved cohort')
        for g in GENOTYPES:
            col=columns[g]
            for row in range(3,45):
                cell=sheet.cell(row,col)
                if cell.value is None: continue
                if cell.data_type=='f' or not isinstance(cell.value,(int,float)):
                    raise ValueError(f'Unexpected measurement {cell.coordinate}')
                rows.append(dict(observation_id=f'12L12D:{cell.coordinate}',plant_id=f'{g}:row{row}',
                    genotype=g,source_row=row,source_cell=cell.coordinate,elapsed_hours=hour,length_mm=float(cell.value)))
    frame=pd.DataFrame(rows)
    assert len(frame)==849 and frame.observation_id.is_unique
    assert set(frame.genotype)==set(GENOTYPES)
    assert np.isfinite(frame.length_mm).all() and (frame.length_mm>0).all()
    eligible=set(frame.loc[frame.elapsed_hours<=36,'plant_id'])
    excluded=frame[~frame.plant_id.isin(eligible)].copy()
    retained=frame[frame.plant_id.isin(eligible)].copy()
    retained['split']=retained.elapsed_hours.map(lambda h:'train' if h<=36 else 'val' if h==48 else 'test')
    plants=retained[['plant_id','genotype','source_row']].drop_duplicates().copy()
    plants['g_idx']=plants.genotype.map(GENOTYPES.index)
    plants=plants.sort_values(['g_idx','source_row']).reset_index(drop=True)
    plants['plant_index']=np.arange(len(plants))
    retained=retained.merge(plants[['plant_id','plant_index']],on='plant_id',validate='many_to_one')
    retained=retained.sort_values(['plant_index','elapsed_hours']).reset_index(drop=True)
    prefix=retained[retained.split.eq('train')]
    config=dict(protocol=PROTOCOL,genotypes=GENOTYPES,source_description='Approved measurement extract preserving source-cell coordinates; empty columns contain no released genotype',source_workbook_sha256=sha(workbook),condition='12L12D',temperature_c=23.,
        observation_unit='longitudinal plant; genotype and source row identify the same plant across times, as confirmed by contributor',
        source_measurements=len(frame),retained_measurements=len(retained),n_plants=len(plants),
        minimum_prefix_observations=1,exclusion_rule='Exclude only plants with no observed length at 0,12,24,36 h; no future-length or future-availability filtering.',
        excluded_observations=excluded.to_dict('records'),prefix_hours=PREFIX_HOURS.tolist(),validation_hours=[48],test_hours=[60,72],
        counts={k:int(len(v)) for k,v in retained.groupby('split')},
        plants_with_observations={k:int(v.plant_id.nunique()) for k,v in retained.groupby('split')},
        prefix_count_distribution={str(k):int(v) for k,v in prefix.groupby('plant_id').size().value_counts().sort_index().items()},
        height_scale_mm=float(prefix.length_mm.max()),scale_fit_on='0--36 h observed training lengths only',
        phenotype_means_used=False,target_type='individual_plant_lengths',
        normalization='training maximum',split_rule='same plants across temporal partitions; no 48/60/72 h phenotype enters the encoder')
    assert config['retained_measurements']==848 and set(plants.genotype)==set(GENOTYPES)
    output=Path(output); output.mkdir(parents=True,exist_ok=True)
    for name,part in [('observations',retained),('plants',plants),('excluded_observations',excluded)]:
        (output/f'{name}.csv').write_text(part.to_csv(index=False))
    for split in ['train','val','test']:
        (output/f'{split}.csv').write_text(retained[retained.split.eq(split)].to_csv(index=False))
    (output/'config.json').write_text(json.dumps(config,indent=2)+'\n')
    return config


def load(device='cpu',include_test=False,grid_hours=HOURS,drop_fraction=0.,mask_seed=20260911):
    config=json.loads((PREPARED/'config.json').read_text())
    plants=pd.read_csv(PREPARED/'plants.csv')
    n=len(plants); grid_hours=np.asarray(grid_hours,dtype=float)
    training=pd.read_csv(PREPARED/'train.csv')
    if not 0<=drop_fraction<1: raise ValueError('drop_fraction must be in [0,1)')
    wanted=int(round(len(training)*drop_fraction)); remaining=training.groupby('plant_index').size().to_dict()
    dropped=[]
    for i in np.random.default_rng(mask_seed).permutation(len(training)):
        if len(dropped)>=wanted: break
        plant=int(training.iloc[i].plant_index)
        if remaining[plant]>1: dropped.append(int(i));remaining[plant]-=1
    if len(dropped)!=wanted: raise ValueError('Cannot drop this many observations while retaining a prefix per plant')
    removed_ids=training.iloc[dropped].observation_id.tolist()
    training=training.drop(index=dropped)
    scale=float(training.length_mm.max())
    config.update(height_scale_mm=scale,additional_prefix_drop=drop_fraction,mask_seed=mask_seed,
        removed_training_observation_ids=removed_ids,retained_training_observations=len(training),
        prefix_missing_fraction=1-len(training)/(4*n))
    assert grid_hours[0]==0 and grid_hours[-1]==72 and np.allclose(np.diff(grid_hours),grid_hours[1]-grid_hours[0])
    prefix=np.zeros((n,4),dtype=np.float32); observed=np.zeros_like(prefix)
    batches={}; raw={}
    for split in (['train','val','test'] if include_test else ['train','val']):
        frame=training if split=='train' else pd.read_csv(PREPARED/f'{split}.csv')
        raw[split]=frame
        idx=np.searchsorted(grid_hours,frame.elapsed_hours.to_numpy())
        np.testing.assert_array_equal(grid_hours[idx],frame.elapsed_hours)
        target=np.zeros((n,len(grid_hours)),dtype=np.float32); mask=np.zeros_like(target)
        target[frame.plant_index,idx]=frame.length_mm.to_numpy()/scale
        mask[frame.plant_index,idx]=1
        batches[split]={'target':torch.tensor(target,device=device),'mask':torch.tensor(mask,device=device)}
        if split=='train':
            pidx=np.searchsorted(PREFIX_HOURS,frame.elapsed_hours.to_numpy())
            np.testing.assert_array_equal(PREFIX_HOURS[pidx],frame.elapsed_hours)
            prefix[frame.plant_index,pidx]=frame.length_mm.to_numpy()/scale
            observed[frame.plant_index,pidx]=1
    assert (observed.sum(1)>0).all()
    light=(grid_hours%24<12).astype(np.float32)
    x=dict(g_idx=torch.tensor(plants.g_idx.to_numpy(),device=device,dtype=torch.long),
        time=torch.tensor(grid_hours/72,device=device,dtype=torch.float32),
        env=torch.tensor(np.broadcast_to(light[None,:,None],(n,len(grid_hours),1)).copy(),device=device),
        prefix_y=torch.tensor(prefix,device=device),prefix_mask=torch.tensor(observed,device=device),
        prefix_time=torch.tensor(PREFIX_HOURS/72,device=device,dtype=torch.float32))
    return config,plants,x,batches,raw


def masked_rmse(prediction,batch):
    target,mask=batch['target'],batch['mask']
    counts=mask.sum(1); eligible=counts>0
    residual=(prediction-target)*mask
    return (residual.square().sum(1)[eligible]/counts[eligible]).clamp_min(1e-12).sqrt().mean()


def score(prediction,batch,scale):
    pred=np.asarray(prediction,dtype=np.float64)*scale
    target=batch['target'].detach().cpu().numpy().astype(np.float64)*scale
    mask=batch['mask'].detach().cpu().numpy().astype(bool)
    counts=mask.sum(1); valid=counts>0
    squared=np.where(mask,(pred-target)**2,0)
    per=np.sqrt(squared.sum(1)[valid]/counts[valid])
    rmse=float(per.mean()); denominator=float(target[mask].mean())
    return dict(rmse=rmse,relative_error=100*rmse/denominator,mean_target=denominator,
        pooled_rmse=float(np.sqrt(squared.sum()/mask.sum())),n_plants=int(valid.sum()),
        n_observations=int(mask.sum()),per_plant_rmse=per.tolist())


if __name__=='__main__':
    import argparse
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook',type=Path)
    parser.add_argument('--output',type=Path,default=PREPARED)
    args=parser.parse_args()
    print(json.dumps(prepare(args.workbook,args.output),indent=2))
