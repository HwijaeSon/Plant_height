"""Verify cohort isolation, paired masks, and every saved prediction score."""
import hashlib
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import numpy as np
import pandas as pd
import torch
from openpyxl import load_workbook
import forecast_data as data
from report_approved_benchmark import ORDER,RESULTS,directory,seeds

ROOT=Path(__file__).resolve().parents[2]

def audit():
    allowed=set(data.GENOTYPES)
    assert allowed=={'Col-0','hy5','MLB','phyAB'}
    wb=load_workbook(data.DATA/'raw/hypocotyl_growth_20230403.xlsx',data_only=False)
    for sheet in wb:
        labels={sheet.cell(2,c).value for c in range(1,sheet.max_column+1)}-{None}
        assert labels==allowed and not sheet._charts and not sheet._images
        for row in sheet:
            for cell in row:
                if cell.row>2 and cell.value is not None:
                    assert sheet.cell(2,cell.column).value in allowed
                    assert isinstance(cell.value,(int,float)) and cell.data_type!='f'
    with TemporaryDirectory() as temp:
        data.prepare(output=Path(temp))
        for file in Path(temp).iterdir():assert file.read_bytes()==(data.PREPARED/file.name).read_bytes()
    cohort=pd.read_csv(data.PREPARED/'observations.csv')
    assert set(cohort.genotype)==allowed and len(cohort)==848 and cohort.plant_id.nunique()==141
    checked=0;checkpoint_count=0;mask_records={}
    for drop,expected_training in [(0.,500),(.25,375),(.5,250)]:
        for seed in [1,2,3]:
            dc,plants,x,batches,raw=data.load(include_test=True,drop_fraction=drop,mask_seed=20260910+seed)
            mask_records[(drop,seed)]=set(dc['removed_training_observation_ids'])
            assert len(raw['train'])==expected_training and len(raw['val'])==116 and len(raw['test'])==232
            assert not set(raw['train'].observation_id)&mask_records[(drop,seed)]
            assert (x['prefix_mask'].sum(1)>0).all()
            assert set(raw['val'].elapsed_hours)=={48} and set(raw['test'].elapsed_hours)=={60,72}
            for model in ORDER:
                if seed not in seeds(drop,model):continue
                folder=directory(drop,model,seed);run=json.loads((folder/'result.json').read_text())
                assert run['status']=='complete' and run['config']['protocol']==data.PROTOCOL
                assert run['config']['n_plants']==141
                assert set(run['config']['data']['genotypes'])==allowed
                assert set(run['config']['data']['removed_training_observation_ids'])==mask_records[(drop,seed)]
                assert run['config']['height_scale_mm']==dc['height_scale_mm']
                with np.load(folder/'predictions.npz',allow_pickle=False) as arrays:
                    assert arrays['prediction'].shape==(141,25) and set(arrays['genotype'])==allowed
                    np.testing.assert_array_equal(arrays['plant_ids'],plants.plant_id.to_numpy(dtype=str))
                    np.testing.assert_array_equal(arrays['prefix_y'],x['prefix_y'].numpy())
                    np.testing.assert_array_equal(arrays['prefix_mask'],x['prefix_mask'].numpy())
                    for split in ['train','val','test']:
                        np.testing.assert_array_equal(arrays[split+'_mask'],batches[split]['mask'].numpy())
                        np.testing.assert_array_equal(arrays[split+'_target'],batches[split]['target'].numpy())
                        score=data.score(arrays['prediction'],batches[split],dc['height_scale_mm'])
                        for metric in ['rmse','relative_error','pooled_rmse','mean_target','n_plants','n_observations']:
                            np.testing.assert_allclose(score[metric],run['metrics'][split][metric],rtol=1e-12,atol=1e-12)
                training_folder=folder.parent if model=='phytoode' else folder
                selection=json.loads((training_folder/'selection.json').read_text())
                assert selection['test_evaluations']==0
                if model not in ['rf','logistic']:
                    checkpoint=training_folder/'checkpoint.pt'
                    digest=hashlib.sha256(checkpoint.read_bytes()).hexdigest()
                    assert digest==selection['checkpoint_sha256']
                    saved=torch.load(checkpoint,map_location='cpu',weights_only=False)
                    assert saved['config']['protocol']==data.PROTOCOL
                    for name,tensor in saved['state_dict'].items():
                        if name in ['g_enc.weight','g_enc.embedding.weight','embedding.weight']:
                            assert tensor.shape[0]==4
                    if model=='phytoode':assert run['checkpoint_sha256']==digest and saved['config']['lambda_ode']==100 and saved['config']['lambda_k']==0
                    checkpoint_count+=1
                checked+=1
    for seed in [1,2,3]:assert mask_records[(0.,seed)]<=mask_records[(.25,seed)]<=mask_records[(.5,seed)]
    assert checked==52 and checkpoint_count==36
    return dict(status='passed',protocol=data.PROTOCOL,genotypes=data.GENOTYPES,n_plants=141,
                retained_observations=848,checked_runs=checked,checked_neural_checkpoints=checkpoint_count,
                raw_workbook_approved_cells_only=True,source_cell_extraction_identical=True,
                masks_shared_nested_and_training_only=True,all_saved_scores_recomputed=True,
                all_checkpoints_use_current_cohort=True)

if __name__=='__main__':
    result=audit()
    output=ROOT/'submission/hypocotyl_validation.json'
    output.write_text(json.dumps(result,indent=2)+'\n')
    print(json.dumps(result,indent=2))
