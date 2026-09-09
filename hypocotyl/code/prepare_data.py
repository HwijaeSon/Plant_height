"""Extract measured cells without inventing longitudinal plant identifiers."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data/raw/hypocotyl_growth_20230403.xlsx"
GENOTYPES = ["Col-0", "hy5", "MLB", "EMS57", "phyAB"]


def main():
    formula = load_workbook(RAW, data_only=False)
    cached = load_workbook(RAW, data_only=True)
    records, checks = [], []
    for sheet in formula:
        for block, hour in enumerate(range(0, 73, 12)):
            first = 2 + 5 * block
            assert sheet.cell(1, first).value == f"{hour}hr"
            for offset, genotype in enumerate(GENOTYPES):
                column = first + offset
                assert sheet.cell(2, column).value == genotype
                values = []
                for row in range(3, 45):
                    cell = sheet.cell(row, column)
                    if cell.value is None:
                        continue
                    if cell.data_type == "f" or not isinstance(cell.value, (int, float)):
                        raise ValueError(f"Unexpected raw value at {sheet.title}:{cell.coordinate}")
                    assert np.isfinite(cell.value) and cell.value > 0
                    values.append(float(cell.value))
                    records.append(dict(observation_id=f"{sheet.title}:{cell.coordinate}",
                        sheet=sheet.title, genotype=genotype, elapsed_hours=hour,
                        length=float(cell.value), source_cell=cell.coordinate,
                        source_row=row, source_column=column))
                expected = cached[sheet.title].cell(45, column).value
                np.testing.assert_allclose(np.mean(values), expected, rtol=1e-12)
                checks.append(dict(sheet=sheet.title, genotype=genotype, elapsed_hours=hour,
                    n=len(values), mean=float(np.mean(values)), sd=float(np.std(values, ddof=1)),
                    spreadsheet_mean=float(expected)))
    out = ROOT / "data/processed"
    out.mkdir(parents=True, exist_ok=True)
    frame = pd.DataFrame(records)
    assert frame.observation_id.is_unique
    frame.to_csv(out / "observations.csv", index=False)
    pd.DataFrame(checks).to_csv(out / "group_statistics.csv", index=False)
    validation = dict(status="passed", source_sha256=hashlib.sha256(RAW.read_bytes()).hexdigest(),
        measured_cells=len(frame), by_sheet=frame.groupby("sheet").size().to_dict(),
        verified_spreadsheet_means=len(checks), raw_rows=[3, 44],
        excluded="Summary means, SDs, t-tests, and plot-input formulas are not observations.",
        identity="Cell identifiers preserve provenance; row positions do not establish repeated plant identity.")
    (out / "extraction_validation.json").write_text(json.dumps(validation, indent=2) + "\n")
    print(json.dumps(validation, indent=2))


if __name__ == "__main__":
    main()
